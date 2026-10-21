"""Loading city coordinates from CSV or from the original Excel workbook."""

from __future__ import annotations

import csv
import re
from dataclasses import dataclass
from typing import List


@dataclass(frozen=True)
class City:
    """A city with 2-D map coordinates (h = horizontal, v = vertical)."""

    label: int
    name: str   # full display name, e.g. "Olympia (Washington)"
    h: float
    v: float


def _full_name(city: str, state: str) -> str:
    state = (state or "").strip()
    city = (city or "").strip()
    return f"{city} ({state})" if state else city


def load_cities(csv_path: str) -> List[City]:
    """Load cities from a CSV with columns: label,city,state,h,v."""
    cities: List[City] = []
    with open(csv_path, newline="", encoding="utf-8") as f:
        reader = csv.DictReader(f)
        for row in reader:
            cities.append(
                City(
                    label=int(row["label"]),
                    name=_full_name(row["city"], row.get("state", "")),
                    h=float(row["h"]),
                    v=float(row["v"]),
                )
            )
    if not cities:
        raise ValueError(f"No cities found in {csv_path}")
    return cities


def load_cities_xlsx(path: str, sheet: str = "MAP") -> List[City]:
    """Load cities directly from the original 'A Walk Through the United States'
    workbook.

    On the ``MAP`` sheet the city table lives in columns B/C/D/E:
    B = "City (State)", C = label, D = H, E = V. The first city appears on the
    header row (row 8), with the rest following, so we scan every row and keep
    those whose label cell holds a number.
    """
    import openpyxl  # imported lazily so CSV use needs no openpyxl

    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb[sheet]
    name_re = re.compile(r"^(.*?)\s*\((.*)\)\s*$")

    cities: List[City] = []
    for r in range(1, ws.max_row + 1):
        name = ws.cell(row=r, column=2).value
        label = ws.cell(row=r, column=3).value
        h = ws.cell(row=r, column=4).value
        v = ws.cell(row=r, column=5).value
        if not isinstance(label, (int, float)) or name is None:
            continue
        if not isinstance(h, (int, float)) or not isinstance(v, (int, float)):
            continue
        m = name_re.match(str(name))
        display = str(name).strip() if not m else _full_name(m.group(1), m.group(2))
        cities.append(City(label=int(label), name=display, h=float(h), v=float(v)))

    if not cities:
        raise ValueError(f"No cities found on sheet '{sheet}' of {path}")
    cities.sort(key=lambda c: c.label)
    return cities
